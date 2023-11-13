import attrs
import sys
# sys.path.append('...')
import nidcpower
from openpyxl import load_workbook, Workbook
import pandas as pd
from config import wiring
import numpy as np
import hightime
import time

@attrs.define
class chnVoltBias:
    description     : str
    resource        : str
    V_force         : str
    I_compl         : str = attrs.field(
        default=1e-3,
        kw_only=True,
    )
    VoltSense       : bool = attrs.field(
        default=False,
        kw_only=True,
    )
    V_compl         : float = attrs.field(
        default=3,
        kw_only=True,
    )
    I_force         : float = attrs.field(
        default=0,
        kw_only=True,
    )


class HCIstress:
    resource: str
    dieCoord: str
    tileName: str
    chnStress: str
    biases: list[chnVoltBias] = attrs.field(
        validator=attrs.validators.deep_iterable(member_validator=attrs.validators.instance_of((chnVoltBias)))
    )

    sourceDelay: float = attrs.field(
        default=1e-3,
        kw_only=True,
    )

    apertureTime: float = attrs.field(
        default=1e-3,
        kw_only=True,
    )

@attrs.define
class CurrentLimit:
    low_level = attrs.field(
        default=10e-6,
        kw_only=True,
    )
    high_level = attrs.field(
        default=1e-3,
        kw_only=True,
    )

@attrs.define
class CurrentLimitRange:
    low_level = attrs.field(
        default=100e-9,
        kw_only=True,
    )
    high_level = attrs.field(
        default=30e-3,
        kw_only=True,
    )


@attrs.define
class Resource:
    device: str
    GATE : list
    DRAIN: list
    SOURCE: list
    BULK: list

@attrs.define
class VtlinSweep:
    resource: Resource
    V_force_D: float
    V_force_S: float
    V_force_B: float
    I_compl_G: float
    I_compl_D: float
    I_compl_S: float
    I_compl_B: float
    V_step_G: float
    V_start_G: float
    V_stop_G: float
    # CurrentLimit: CurrentLimit
    # CurrentLimitRange: CurrentLimitRange
    sourceDelay: float = attrs.field(
        default=1e-3,
        kw_only=True,
    )
    apertureTime: float = attrs.field(
        default=1e-3,
        kw_only=True,
    )




@attrs.define
class VtSatSweep:
    resource: Resource
    V_force_D: float
    V_force_S: float
    V_force_B: float
    I_compl_G: float
    I_compl_D: float
    I_compl_S: float
    I_compl_B: float
    V_step_G: float
    V_start_G: float
    V_stop_G: float
    CurrentLimit: CurrentLimit               =  attrs.field(
        default=CurrentLimit(),
        kw_only=True,
    )

    CurrentLimitRange: CurrentLimitRange     = attrs.field(
        default=CurrentLimitRange(),
        kw_only=True,
    )


@attrs.define
class IbMaxSweep:
    resource: Resource
    V_force_D: float
    V_force_S: float
    V_force_B: float
    V_force_G: float
    I_compl_G: float
    I_compl_D: float
    I_compl_S: float
    I_compl_B: float
    V_step_G: float
    V_start_G: float
    V_stop_G: float
    CurrentLimit: CurrentLimit               =  attrs.field(
        default=CurrentLimit(),
        kw_only=True,
    )

    CurrentLimitRange: CurrentLimitRange     = attrs.field(
        default=CurrentLimitRange(),
        kw_only=True,
    )



@attrs.define
class HCIsweep:
    resource: Resource
    description: str
    Vtlin: VtlinSweep
    VtSat: VtSatSweep
    IbMax: IbMaxSweep


    dieCoord: str = attrs.field(
        default='A1',
        kw_only=True,
    )

    tileName: str = attrs.field(
        default='A1',
        kw_only=True,
    )


def dumpExceltoPython(testplan, sitesNum):
    wb = load_workbook(filename=testplan)
    ws = wb.active
    dDut, dHCItest, dTest_Vtlin, dTest_VtSat, dTest_IbMAX, dResource = {}, {}, {}, {}, {}, {}
    header = []
    rows_buf = []
    count = 0
    for row in ws.iter_rows(min_row=None, max_col=None, max_row=None, values_only=True):
        if count == 0:
            header = row
            # print(header)
        else:
            rows_buf.append(row)
        count += 1
    df_plan = pd.DataFrame(rows_buf, columns=header)
    norepeat_df_plan = df_plan.drop_duplicates(subset=[' dut.name']).reset_index()
    # count_1 = 0
    sDuts = set()

    # Get all the DUTs and resources
    count = 0
    for dut in norepeat_df_plan.loc[:, ' dut.name']:
        # print(norepeat_df_plan.loc[:, ' dut.pads[G]'])
        padGATE = norepeat_df_plan.loc[count, ' dut.pads[G]']
        # print(f'Gate is {padGATE}')
        padDRAIN = norepeat_df_plan.loc[count, ' dut.pads[D]']
        padSOURCE = norepeat_df_plan.loc[count, ' dut.pads[S]']
        padBULK = norepeat_df_plan.loc[count, ' dut.pads[B]']
        GATE = [f'{wiring[(padGATE-1)+24*i][1][0][wiring[(padGATE-1)+24*i][1][1]]}/{wiring[(padGATE-1)+24*i][1][2]}' for i in range(sitesNum)]
        DRAIN  = [f'{wiring[(padDRAIN-1)+24*i][1][0][wiring[(padDRAIN-1)+24*i][1][1]]}/{wiring[(padDRAIN-1)+24*i][1][2]}' for i in range(sitesNum)]
        SOURCE = [f'{wiring[(padSOURCE-1)+24*i][1][0][wiring[(padSOURCE-1)+24*i][1][1]]}/{wiring[(padSOURCE-1)+24*i][1][2]}' for i in range(sitesNum)]
        BULK = [f'{wiring[(padBULK-1)+24*i][1][0][wiring[(padBULK-1)+24*i][1][1]]}/{wiring[(padBULK-1)+24*i][1][2]}' for i in range(sitesNum)]
        dResource[dut] = Resource(dut, GATE, DRAIN, SOURCE, BULK)
        sDuts.add(dut)
        count += 1
    # print(dResource)

    count = 0
    for c in df_plan.loc[:, 'fullname']:
        dut = df_plan.loc[count, ' dut.name']
        if 'Vtlin' in c:
            vForceD = df_plan.loc[count,  ' test.V_force_D']
            vForceS = df_plan.loc[count,  ' test.V_force_S']
            vForceB = df_plan.loc[count,  ' test.V_force_B']
            vStartG = df_plan.loc[count,  ' test.V_start_G']
            vStopG = df_plan.loc[count,  ' test.V_stop_G']
            vStepG = df_plan.loc[count,  ' test.V_step_G'] 
            iComplG = df_plan.loc[count,  ' test.I_compl_G']
            iComplD = df_plan.loc[count,  ' test.I_compl_D']
            iComplS = df_plan.loc[count,  ' test.I_compl_S']
            iComplB = df_plan.loc[count,  ' test.I_compl_B']
            dTest_Vtlin[dut] = VtlinSweep(dResource[dut], vForceD, vForceS, vForceB, iComplG, iComplD, iComplS, iComplB, vStepG, vStartG, vStopG)
        elif 'Vtsat' in c:
            vForceD = df_plan.loc[count,  ' test.V_force_D']
            vForceS = df_plan.loc[count,  ' test.V_force_S']
            vForceB = df_plan.loc[count,  ' test.V_force_B']
            vStartG = df_plan.loc[count,  ' test.V_start_G']
            vStopG = df_plan.loc[count,  ' test.V_stop_G']
            vStepG = df_plan.loc[count,  ' test.V_step_G'] 
            iComplG = df_plan.loc[count,  ' test.I_compl_G']
            iComplD = df_plan.loc[count,  ' test.I_compl_D']
            iComplS = df_plan.loc[count,  ' test.I_compl_S']
            iComplB = df_plan.loc[count,  ' test.I_compl_B']
            dTest_VtSat[dut] = VtSatSweep(dResource[dut], vForceD, vForceS, vForceB, iComplG, iComplD, iComplS, iComplB, vStepG, vStartG, vStopG)
        elif 'Ibmax' in c:
            vForceD = df_plan.loc[count,  ' test.V_force_D'] # stress Drain voltage
            vForceS = df_plan.loc[count,  ' test.V_force_S']
            vForceB = df_plan.loc[count,  ' test.V_force_B']
            vStartG = df_plan.loc[count,  ' test.V_start_G']
            vStopG = df_plan.loc[count,  ' test.V_stop_G']
            vStepG = df_plan.loc[count,  ' test.V_step_G']
            vForceG = df_plan.loc[count,  ' test.V_force_G'] # stress Gate voltage
            iComplG = df_plan.loc[count,  ' test.I_compl_G']
            iComplD = df_plan.loc[count,  ' test.I_compl_D']
            iComplS = df_plan.loc[count,  ' test.I_compl_S']
            iComplB = df_plan.loc[count,  ' test.I_compl_B']
            dTest_IbMAX[dut] = IbMaxSweep(dResource[dut], vForceD, vForceS, vForceB, iComplG, iComplD, iComplS, iComplB, iComplG, vStepG, vStartG, vStopG)
        # print(dTest[c])
        count += 1

    # print("hahha")
    for dut in sDuts:
        resource = dResource[dut]
        Vtlin = dTest_Vtlin[dut]
        VtSat = dTest_VtSat[dut]
        Ibmax = dTest_IbMAX[dut]
        dHCItest[dut] = HCIsweep(dResource[dut], dut, dTest_Vtlin[dut], dTest_VtSat[dut], dTest_IbMAX[dut])
    for k, v in dHCItest.items():
        print(k, v)
    return dHCItest

dumpExceltoPython('test_plan_HCI_1110_V3.xlsx', sitesNum=2)